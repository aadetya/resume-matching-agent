from __future__ import annotations

import json
import os
from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

import fs_tools
from fs_tools import FileToolError, list_files, read_file, search_in_file, write_file


def test_read_utf8_txt_returns_content_metadata_and_segment(tmp_path: Path) -> None:
    path = tmp_path / "resume.txt"
    path.write_text("Ada Example\nPython engineer\n", encoding="utf-8")

    result = read_file(str(path))

    assert result["ok"] is True
    assert result["content"] == "Ada Example\nPython engineer"
    assert result["metadata"]["encoding"] == "utf-8"
    assert result["metadata"]["word_count"] == 4
    assert len(result["metadata"]["sha256"]) == 64
    assert result["segments"][0]["label"] == "document"
    json.dumps(result)


def test_read_non_utf8_txt_detects_encoding(tmp_path: Path) -> None:
    path = tmp_path / "latin.txt"
    path.write_bytes("Renée Example\nPython".encode("cp1252"))

    result = read_file(str(path))

    assert result["ok"] is True
    assert "Ren" in result["content"]
    assert result["warnings"]


def test_read_rejects_binary_txt(tmp_path: Path) -> None:
    path = tmp_path / "binary.txt"
    path.write_bytes(b"hello\x00world")

    result = read_file(str(path))

    assert result["ok"] is False
    assert result["error"]["code"] == "binary_text_file"


def test_read_reports_missing_and_unsupported_files(tmp_path: Path) -> None:
    missing = read_file(str(tmp_path / "missing.pdf"))
    unsupported_path = tmp_path / "resume.rtf"
    unsupported_path.write_text("text", encoding="utf-8")
    unsupported = read_file(str(unsupported_path))

    assert missing["error"]["code"] == "not_found"
    assert unsupported["error"]["code"] == "unsupported_format"


def test_read_rejects_extension_signature_mismatch(tmp_path: Path) -> None:
    path = tmp_path / "pretend.pdf"
    path.write_text("not a PDF", encoding="utf-8")

    result = read_file(str(path))

    assert result["ok"] is False
    assert result["error"]["code"] == "format_mismatch"


def test_read_rejects_corrupt_docx(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"PK-not-really-a-zip")

    result = read_file(str(path))

    assert result["ok"] is False
    assert result["error"]["code"] == "corrupt_document"


def test_read_rejects_encrypted_pdf(tmp_path: Path) -> None:
    path = tmp_path / "locked.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("secret")
    with path.open("wb") as stream:
        writer.write(stream)

    result = read_file(str(path))

    assert result["ok"] is False
    assert result["error"]["code"] == "encrypted_document"


def test_read_applies_file_size_limit(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "resume.txt"
    path.write_text("too large", encoding="utf-8")
    monkeypatch.setattr(fs_tools, "MAX_FILE_BYTES", 2)

    result = read_file(str(path))

    assert result["ok"] is False
    assert result["error"]["code"] == "file_too_large"


def test_docx_extraction_preserves_paragraph_table_order(tmp_path: Path) -> None:
    path = tmp_path / "ordered.docx"
    document = Document()
    document.add_paragraph("Before table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "advanced"
    document.add_paragraph("After table")
    document.save(str(path))

    result = read_file(str(path))

    assert result["ok"] is True
    assert result["content"].index("Before table") < result["content"].index("Python")
    assert result["content"].index("Python") < result["content"].index("After table")
    kinds = [segment["kind"] for segment in result["segments"]]
    assert kinds == ["paragraph", "table_row", "paragraph"]
    assert result["metadata"]["table_count"] == 1


def test_pdf_extraction_and_page_provenance(tmp_path: Path) -> None:
    path = tmp_path / "two-pages.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 720, "First page: SQL")
    pdf.showPage()
    pdf.drawString(72, 720, "Second page: PYTHON delivery")
    pdf.save()

    read_result = read_file(str(path))
    search_result = search_in_file(str(path), "python")

    assert read_result["ok"] is True
    assert read_result["metadata"]["page_count"] == 2
    assert search_result["match_count"] == 1
    assert search_result["matches"][0]["location"]["label"] == "page 2"


def test_list_files_filters_case_insensitively_and_sorts(tmp_path: Path) -> None:
    (tmp_path / "z.PDF").write_bytes(b"%PDF-")
    (tmp_path / "A.pdf").write_bytes(b"%PDF-")
    (tmp_path / "note.txt").write_text("x", encoding="utf-8")
    (tmp_path / "nested").mkdir()

    files = list_files(str(tmp_path), "pdf")

    assert [item["name"] for item in files] == ["A.pdf", "z.PDF"]
    assert all(item["extension"] == ".pdf" for item in files)
    assert all(item["modified_at"].endswith("Z") for item in files)


def test_list_files_raises_typed_error_for_invalid_directory(tmp_path: Path) -> None:
    with pytest.raises(FileToolError) as captured:
        list_files(str(tmp_path / "absent"))
    assert captured.value.code == "not_found"


def test_write_file_creates_directories_and_overwrites_atomically(tmp_path: Path) -> None:
    path = tmp_path / "nested" / "summary.md"

    created = write_file(str(path), "first")
    overwritten = write_file(str(path), "second")

    assert created["ok"] is True and created["status"] == "created"
    assert overwritten["ok"] is True and overwritten["status"] == "overwritten"
    assert path.read_text(encoding="utf-8") == "second"
    assert not list(path.parent.glob(".*.tmp"))
    assert overwritten["metadata"]["character_count"] == 6


def test_write_file_rejects_non_string_content(tmp_path: Path) -> None:
    result = write_file(str(tmp_path / "summary.md"), 123)  # type: ignore[arg-type]

    assert result["ok"] is False
    assert result["error"]["code"] == "invalid_argument"


def test_search_is_literal_case_insensitive_and_returns_context(tmp_path: Path) -> None:
    path = tmp_path / "resume.txt"
    path.write_text("Heading\nBuilt C++ services\nThen wrote c++ tests\nFooter", encoding="utf-8")

    result = search_in_file(str(path), "C++")

    assert result["ok"] is True
    assert result["match_count"] == 2
    assert result["matches"][0]["start_line"] == 2
    assert "Heading" in result["matches"][0]["context"]
    assert "Then wrote" in result["matches"][0]["context"]


def test_search_docx_match_reports_table_row(tmp_path: Path) -> None:
    path = tmp_path / "resume.docx"
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "Python"
    document.save(str(path))

    result = search_in_file(str(path), "python")

    assert result["ok"] is True
    assert result["matches"][0]["location"]["kind"] == "table_row"
    assert result["matches"][0]["location"]["table_index"] == 1


def test_search_limits_returned_matches(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "resume.txt"
    path.write_text("Python " * 12, encoding="utf-8")
    monkeypatch.setattr(fs_tools, "MAX_SEARCH_MATCHES", 5)

    result = search_in_file(str(path), "python")

    assert result["match_count"] == 12
    assert result["returned_match_count"] == 5
    assert result["truncated"] is True


def test_search_counts_dense_matches_and_bounds_each_context(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "dense.txt"
    path.write_text("x" * 50 + "a" * 200_000, encoding="utf-8")
    monkeypatch.setattr(fs_tools, "MAX_SEARCH_MATCHES", 5)
    monkeypatch.setattr(fs_tools, "SEARCH_CONTEXT_CHARS", 20)

    result = search_in_file(str(path), "a")

    assert result["match_count"] == 200_000
    assert result["returned_match_count"] == 5
    assert result["truncated"] is True
    assert all(len(match["context"]) <= 43 for match in result["matches"])
    assert sum(len(match["context"]) for match in result["matches"]) <= 215
    assert result["matches"][0]["context"].startswith("…")
    assert result["matches"][0]["context"].endswith("…")


def test_search_rejects_empty_keyword(tmp_path: Path) -> None:
    path = tmp_path / "resume.txt"
    path.write_text("Python", encoding="utf-8")

    result = search_in_file(str(path), "")

    assert result["ok"] is False
    assert result["error"]["code"] == "invalid_keyword"


@pytest.mark.skipif(os.name == "nt", reason="symlink permission semantics differ on Windows")
def test_list_metadata_marks_symlink(tmp_path: Path) -> None:
    target = tmp_path / "real.txt"
    link = tmp_path / "linked.txt"
    target.write_text("data", encoding="utf-8")
    link.symlink_to(target)

    records = {item["name"]: item for item in list_files(str(tmp_path), ".txt")}

    assert records["linked.txt"]["is_symlink"] is True
    assert records["linked.txt"]["path"] == str(link.absolute())
    assert records["linked.txt"]["path"] != str(target.resolve())
    assert records["linked.txt"]["size_bytes"] == link.lstat().st_size
    assert records["linked.txt"]["readable"] is None
