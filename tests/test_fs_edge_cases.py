from __future__ import annotations

import codecs
import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter

from screening_agent import file_tools


@pytest.mark.parametrize(
    "bom,encoding",
    [
        (codecs.BOM_UTF8, "utf-8"),
        (codecs.BOM_UTF16_LE, "utf-16-le"),
        (codecs.BOM_UTF16_BE, "utf-16-be"),
        (codecs.BOM_UTF32_LE, "utf-32-le"),
        (codecs.BOM_UTF32_BE, "utf-32-be"),
    ],
)
def test_unicode_bom_decodes_before_binary_rejection(tmp_path, bom, encoding):
    path = tmp_path / "resume.txt"
    path.write_bytes(bom + "José Pérez\nReact developer".encode(encoding))
    result = file_tools.read_file(str(path))
    assert result["ok"]
    assert result["content"] == "José Pérez\nReact developer"


def test_malformed_unicode_bom_is_typed_error(tmp_path):
    path = tmp_path / "resume.txt"
    path.write_bytes(codecs.BOM_UTF16_LE + b"X")
    result = file_tools.read_file(str(path))
    assert result["error"]["code"] == "encoding_error"


def test_decoded_binary_control_bytes_are_rejected(tmp_path):
    path = tmp_path / "resume.txt"
    path.write_bytes(b"normal prefix\x03\x08payload")
    assert file_tools.read_file(str(path))["error"]["code"] == "binary_text_file"


def test_truncated_pdf_has_corrupt_document_result(tmp_path):
    path = tmp_path / "resume.pdf"
    path.write_bytes(b"%PDF-1.7\n1 0 obj << /Type /Catalog >> endobj\n")
    assert file_tools.read_file(str(path))["error"]["code"] == "corrupt_document"


def test_blank_pdf_explicitly_reports_ocr_requirement(tmp_path):
    path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with path.open("wb") as stream:
        writer.write(stream)
    result = file_tools.read_file(str(path))
    assert result["ok"] and result["content"] == ""
    assert result["metadata"]["ocr_required"] is True


def test_docx_zip_with_broken_document_xml_is_corrupt(tmp_path):
    original = io.BytesIO()
    Document().save(original)
    path = tmp_path / "broken.docx"
    with zipfile.ZipFile(io.BytesIO(original.getvalue())) as old, zipfile.ZipFile(path, "w") as new:
        for name in old.namelist():
            new.writestr(name, b"<not closed" if name == "word/document.xml" else old.read(name))
    assert file_tools.read_file(str(path))["error"]["code"] == "corrupt_document"


def test_docx_decompression_limit_checked_before_parser(tmp_path, monkeypatch):
    path = tmp_path / "overlarge.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "a" * 100)
        archive.writestr("word/document.xml", "a" * 100)
    monkeypatch.setattr(file_tools, "MAX_DOCX_UNCOMPRESSED_BYTES", 10)
    assert file_tools.read_file(str(path))["error"]["code"] == "unsafe_document"


def test_docx_header_is_preserved(tmp_path):
    path = tmp_path / "header.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Jane Doe"
    document.add_paragraph("Engineer | Example Systems | Jan 2020 - Jan 2024")
    document.save(path)
    result = file_tools.read_file(str(path))
    assert result["content"].startswith("Jane Doe")
    assert result["segments"][0]["kind"] == "header"


def test_pdf_late_page_failure_is_corrupt(tmp_path, monkeypatch):
    path = tmp_path / "late.pdf"
    path.write_bytes(b"%PDF-1.7")

    class Page:
        def extract_text(self):
            raise ValueError("damaged page stream")

    class Reader:
        is_encrypted = False
        pages = [Page()]

    monkeypatch.setattr(file_tools, "PdfReader", lambda *args, **kwargs: Reader())
    assert file_tools.read_file(str(path))["error"]["code"] == "corrupt_document"


def test_nested_docx_table_preserves_cell_order(tmp_path):
    path = tmp_path / "nested.docx"
    document = Document()
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].text = "Before nested"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "React nested evidence"
    cell.add_paragraph("After nested")
    document.save(path)
    result = file_tools.read_file(str(path))
    assert result["ok"]
    assert (
        result["content"].index("Before nested")
        < result["content"].index("React nested evidence")
        < result["content"].index("After nested")
    )


def test_docx_late_extraction_error_has_typed_result(tmp_path, monkeypatch):
    path = tmp_path / "late.docx"
    Document().save(path)

    def broken_blocks(document):
        raise ValueError("damaged inner document")

    monkeypatch.setattr(file_tools, "_iter_docx_blocks", broken_blocks)
    assert file_tools.read_file(str(path))["error"]["code"] == "corrupt_document"
